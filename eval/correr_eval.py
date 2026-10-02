#!/usr/bin/env python3
"""Corre la evaluación de modelos locales contra el backend (POST /comparar) y llena el Excel.

Uso (desde la carpeta del proyecto):
  eval/.venv/bin/python eval/correr_eval.py
      Corrida completa: 7 modelos × 25 casos, guarda resultados y llena la pestaña Casos.
  eval/.venv/bin/python eval/correr_eval.py --casos M01 --modelos mistral:latest --sin-excel
      Prueba rápida sin tocar el Excel.
  eval/.venv/bin/python eval/correr_eval.py --desde eval/resultados/<fecha-hora>.jsonl
      No llama al backend: llena el Excel con una corrida ya guardada.
"""
import argparse
import json
import logging
import re
import shutil
import sys
import time
import warnings
import xml.etree.ElementTree as ET
import zipfile
from collections import defaultdict
from datetime import datetime
from pathlib import Path

import httpx
import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation

EVAL = Path(__file__).resolve().parent
PROYECTO = EVAL.parent
EXCEL = PROYECTO / "modelos-ollama.xlsx"
CASOS = EVAL / "casos_eval.jsonl"
PROMPTS = EVAL / "prompts_para_comparador.md"
RESULTADOS = EVAL / "resultados"

BACKEND = "http://localhost:8000"
OLLAMA = "http://localhost:11434"
TIMEOUT_S = 300
MAX_TOKENS = 512

MARCADOR_OCR = "<<< PEGA AQUÍ LA SALIDA DE glm-ocr >>>"
ANSI = re.compile(r"\x1b\[[0-9;]*[A-Za-z]")
MSG_SIN_ESQUEMA = "No encontré el esquema"

FILA_INI, FILA_FIN = 7, 31          # Casos!A7:A31
CELDAS_MODELOS = range(5, 12)       # Evaluación!A5:A11
LISTA_SI_NO = '"Sí,No"'

NS = {
    "m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
    "x14": "http://schemas.microsoft.com/office/spreadsheetml/2009/9/main",
    "xm": "http://schemas.microsoft.com/office/excel/2006/main",
}

log = logging.getLogger("eval")


class SinEsquema(Exception):
    pass


# ---------------------------------------------------------------- entradas

def cargar_casos():
    with CASOS.open(encoding="utf-8") as f:
        return [json.loads(linea) for linea in f if linea.strip()]


def limpiar_ocr(ruta):
    """Primera copia del OCR (antes de la primera línea con ```), sin secuencias ANSI."""
    primera = []
    for linea in ruta.read_text(encoding="utf-8").splitlines():
        if linea.startswith("```"):
            break
        primera.append(linea)
    return ANSI.sub("", "\n".join(primera)).strip()


def cargar_prompts(casos):
    md = PROMPTS.read_text(encoding="utf-8")
    mensajes = {
        m.group(1): m.group(2)
        for m in re.finditer(r"^### (M\d{2}) · .*?\n```text\n(.*?)\n```", md, re.S | re.M)
    }
    plantilla = re.search(r"```text\n(.*?)\n```", md.split("## Documentos", 1)[1], re.S).group(1)
    if MARCADOR_OCR not in plantilla:
        sys.exit(f"La plantilla de documentos no trae el marcador {MARCADOR_OCR!r}")

    prompts = {}
    for caso in casos:
        cid = caso["id"]
        if cid.startswith("M"):
            if cid not in mensajes:
                sys.exit(f"No encontré el bloque ```text de {cid} en {PROMPTS.name}")
            prompts[cid] = mensajes[cid]
        else:
            prompts[cid] = plantilla.replace(MARCADOR_OCR, limpiar_ocr(EVAL / f"ocr_{cid}.txt"))
    return prompts


def leer_modelos(wb):
    ws = wb["Evaluación"]
    return [(ws[f"A{r}"].value or "").strip() or None for r in CELDAS_MODELOS]


def modelos_instalados(cliente):
    r = cliente.get(f"{OLLAMA}/api/tags")
    r.raise_for_status()
    nombres = set()
    for m in r.json()["models"]:
        nombres.add(m["name"])
        if m["name"].endswith(":latest"):
            nombres.add(m["name"].removesuffix(":latest"))
    return nombres


# ---------------------------------------------------------------- backend

def armar_request(caso, modelo, prompt):
    return {
        "prompt": prompt,
        "modelos": [modelo],
        "esquema": "mensaje" if caso["id"].startswith("M") else "documento",
        "secuencial": True,
        "esperado": caso["esperado"],
        "max_tokens": MAX_TOKENS,
    }


def llamar_backend(cliente, body):
    """Devuelve (respuesta, error_cliente). Lanza SinEsquema si el backend no encuentra esquemas.json."""
    try:
        r = cliente.post(f"{BACKEND}/comparar", json=body)
    except httpx.HTTPError as e:
        return None, f"{type(e).__name__}: {e}"
    try:
        respuesta = r.json()
    except ValueError:
        respuesta = {"_cuerpo": r.text}
    if MSG_SIN_ESQUEMA in r.text or MSG_SIN_ESQUEMA in json.dumps(respuesta, ensure_ascii=False):
        raise SinEsquema(r.text[:500])
    if r.status_code != 200:
        return respuesta, f"HTTP {r.status_code}: {r.text[:300]}"
    return respuesta, None


def resumir(caso, modelo, respuesta, error_cliente):
    """Lo que va al Excel y a la tabla: json_valido, puntaje, n, latencia (s) y error."""
    fila = {"caso": caso["id"], "modelo": modelo, "n": caso["n_campos"],
            "json_valido": False, "puntaje": 0, "latencia": None, "error": error_cliente}
    if error_cliente:
        return fila

    resultados = (respuesta or {}).get("resultados") or []
    res = next((x for x in resultados if x.get("modelo") == modelo), resultados[0] if resultados else None)
    if res is None:
        fila["error"] = "La respuesta no trae resultados"
        return fila
    if res.get("error"):
        fila["error"] = str(res["error"])
        return fila

    ev = res.get("evaluacion") or {}
    if ev.get("n_esperados") not in (None, caso["n_campos"]):
        log.warning("%s · %s: n_esperados=%s del backend ≠ n_campos=%s del caso",
                    modelo, caso["id"], ev["n_esperados"], caso["n_campos"])
    puntaje = ev.get("puntaje") or 0
    fila["puntaje"] = int(puntaje) if float(puntaje).is_integer() else puntaje
    fila["json_valido"] = bool(res.get("json_valido"))
    t = res.get("tiempos_ollama")
    if t:
        fila["latencia"] = round(((t.get("prompt_ms") or 0) + (t.get("generacion_ms") or 0)) / 1000, 1)
    return fila


def correr(casos, prompts, modelos, salida):
    """Por modelo y, dentro de cada modelo, todos los casos (el modelo se queda cargado)."""
    filas, omitidos = [], {}
    with httpx.Client(timeout=TIMEOUT_S) as cliente, salida.open("w", encoding="utf-8") as f:
        instalados = modelos_instalados(cliente)
        for modelo in modelos:
            if modelo not in instalados:
                log.warning("%s no está instalado en Ollama; lo salto", modelo)
                omitidos[modelo] = "NO INSTALADO"
                f.write(json.dumps({"caso": None, "modelo": modelo, "omitido": "no instalado en Ollama"},
                                   ensure_ascii=False) + "\n")
                f.flush()
                continue
            for caso in casos:
                body = armar_request(caso, modelo, prompts[caso["id"]])
                t0 = time.monotonic()
                respuesta, error_cliente = llamar_backend(cliente, body)
                f.write(json.dumps({"caso": caso["id"], "modelo": modelo, "request": body,
                                    "respuesta": respuesta, "error_cliente": error_cliente},
                                   ensure_ascii=False) + "\n")
                f.flush()
                fila = resumir(caso, modelo, respuesta, error_cliente)
                filas.append(fila)
                if fila["error"]:
                    log.error("%s · %s: %s", modelo, caso["id"], fila["error"])
                else:
                    log.info("%s · %s: json=%s puntaje=%s/%s lat=%ss (pared %.1fs)",
                             modelo, caso["id"], "Sí" if fila["json_valido"] else "No",
                             fila["puntaje"], fila["n"], fila["latencia"], time.monotonic() - t0)
    return filas, omitidos


def cargar_corrida(ruta, casos):
    por_id = {c["id"]: c for c in casos}
    filas, omitidos = [], {}
    with ruta.open(encoding="utf-8") as f:
        for linea in f:
            d = json.loads(linea)
            if d.get("omitido"):
                omitidos[d["modelo"]] = "NO INSTALADO"
                continue
            filas.append(resumir(por_id[d["caso"]], d["modelo"], d["respuesta"], d.get("error_cliente")))
    return filas, omitidos


# ---------------------------------------------------------------- Excel

def abrir_excel(ruta):
    """Carga con openpyxl y dice si avisó que quita validaciones de datos."""
    with warnings.catch_warnings(record=True) as avisos:
        warnings.simplefilter("always")
        wb = openpyxl.load_workbook(ruta)
    quito = any("Data Validation" in str(a.message) for a in avisos)
    return wb, quito


def validaciones_x14(ruta):
    """Validaciones de la extensión x14 (listas que apuntan a otra hoja) que openpyxl descarta, por hoja."""
    with zipfile.ZipFile(ruta) as z:
        libro = ET.fromstring(z.read("xl/workbook.xml"))
        rels = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
        destino = {r.get("Id"): r.get("Target") for r in rels.findall("rel:Relationship", NS)}
        por_hoja = defaultdict(list)
        for hoja in libro.find("m:sheets", NS):
            target = destino[hoja.get(f"{{{NS['r']}}}id")].lstrip("/")
            parte = target if target.startswith("xl/") else f"xl/{target}"
            for dv in ET.fromstring(z.read(parte)).iter(f"{{{NS['x14']}}}dataValidation"):
                f1 = dv.find("x14:formula1/xm:f", NS)
                f2 = dv.find("x14:formula2/xm:f", NS)
                por_hoja[hoja.get("name")].append({
                    "type": dv.get("type"),
                    "operator": dv.get("operator"),
                    "formula1": f1.text if f1 is not None else None,
                    "formula2": f2.text if f2 is not None and dv.get("type") != "list" else None,
                    "allowBlank": dv.get("allowBlank") in ("1", "true"),
                    "showDropDown": dv.get("showDropDown") in ("1", "true"),
                    "showInputMessage": dv.get("showInputMessage") in ("1", "true"),
                    "showErrorMessage": dv.get("showErrorMessage") in ("1", "true"),
                    "errorTitle": dv.get("errorTitle"), "error": dv.get("error"),
                    "promptTitle": dv.get("promptTitle"), "prompt": dv.get("prompt"),
                    "sqref": dv.find("xm:sqref", NS).text,
                })
    return por_hoja


def reponer_validaciones(wb, x14, n_modelos):
    # Lista "Sí,No" en las columnas de JSON válido de Casos (G, J, M, ... filas 7–31).
    ws = wb["Casos"]
    rangos = {f"{openpyxl.utils.get_column_letter(7 + 3 * k)}{FILA_INI}:"
              f"{openpyxl.utils.get_column_letter(7 + 3 * k)}{FILA_FIN}" for k in range(n_modelos)}
    cubiertos = {str(r) for dv in ws.data_validations.dataValidation
                 if dv.type == "list" and dv.formula1 == LISTA_SI_NO for r in dv.sqref.ranges}
    if not rangos <= cubiertos:
        ws.add_data_validation(DataValidation(type="list", formula1=LISTA_SI_NO, sqref=" ".join(sorted(rangos))))
        log.info("Repuse la lista Sí,No en Casos: %s", " ".join(sorted(rangos)))
    # Lo que openpyxl descartó del archivo original, en la hoja donde estaba.
    for hoja, dvs in x14.items():
        for d in dvs:
            wb[hoja].add_data_validation(DataValidation(**d))
            log.info("Repuse la validación %s %s en %s!%s", d["type"], d["formula1"], hoja, d["sqref"])


def firma_validaciones(ws):
    return {(dv.type, dv.formula1, frozenset(str(dv.sqref).split())) for dv in ws.data_validations.dataValidation}


def verificar(antes, despues, x14, columnas_editables):
    """Compara el respaldo con el archivo nuevo: solo pueden cambiar las celdas de resultados de Casos."""
    wa, _ = abrir_excel(antes)
    wd, _ = abrir_excel(despues)
    problemas = []
    for nombre in wa.sheetnames:
        a, d = wa[nombre], wd[nombre]
        for r in range(1, max(a.max_row, d.max_row) + 1):
            for c in range(1, max(a.max_column, d.max_column) + 1):
                if nombre == "Casos" and FILA_INI <= r <= FILA_FIN and c in columnas_editables:
                    continue
                if a.cell(r, c).value != d.cell(r, c).value:
                    problemas.append(f"{nombre}!{a.cell(r, c).coordinate}: {a.cell(r, c).value!r} → {d.cell(r, c).value!r}")
                if bool(a.cell(r, c).comment) != bool(d.cell(r, c).comment):
                    problemas.append(f"{nombre}!{a.cell(r, c).coordinate}: cambió el comentario")
        esperadas = firma_validaciones(a) | {(v["type"], v["formula1"], frozenset(v["sqref"].split()))
                                             for v in x14.get(nombre, [])}
        if firma_validaciones(d) != esperadas:
            problemas.append(f"{nombre}: validaciones distintas {firma_validaciones(d) ^ esperadas}")
        if set(map(str, a.merged_cells.ranges)) != set(map(str, d.merged_cells.ranges)):
            problemas.append(f"{nombre}: celdas combinadas distintas")
    return problemas


def escribir_excel(filas, omitidos, modelos, stamp):
    lock = EXCEL.with_name(f"~${EXCEL.name}")
    if lock.exists():
        sys.exit(f"El Excel parece estar abierto ({lock.name}). Ciérralo y corre con --desde <jsonl>.")

    respaldo = EXCEL.with_name(f"{EXCEL.stem}.respaldo-{stamp}{EXCEL.suffix}")
    shutil.copy2(EXCEL, respaldo)
    log.info("Respaldo: %s", respaldo.name)

    wb, quito = abrir_excel(EXCEL)
    ws = wb["Casos"]
    fila_de = {ws.cell(r, 1).value: r for r in range(FILA_INI, FILA_FIN + 1)}
    por_par = {(f["modelo"], f["caso"]): f for f in filas}
    casos_corridos = {f["caso"] for f in filas} or set(fila_de)
    columnas_editables = set()

    for k, modelo in enumerate(modelos):
        if modelo is None or (modelo not in omitidos and not any(f["modelo"] == modelo for f in filas)):
            continue  # modelo fuera de esta corrida: no se toca
        c_json, c_ok, c_lat = 7 + 3 * k, 8 + 3 * k, 9 + 3 * k
        columnas_editables |= {c_json, c_ok, c_lat}
        for caso_id, r in fila_de.items():
            if caso_id not in casos_corridos:
                continue
            f = por_par.get((modelo, caso_id))
            if f is None:                      # modelo omitido: se borran los datos viejos
                valores = (None, None, None)
            elif f["error"]:
                valores = ("No", 0, None)
            else:
                valores = ("Sí" if f["json_valido"] else "No", f["puntaje"], f["latencia"])
            for c, v in zip((c_json, c_ok, c_lat), valores):
                ws.cell(r, c).value = v

    x14 = validaciones_x14(EXCEL) if quito else {}
    if quito:
        log.warning("openpyxl avisó que quita validaciones de datos; las repongo")
        reponer_validaciones(wb, x14, len(modelos))

    wb.calculation.fullCalcOnLoad = True
    wb.save(EXCEL)
    log.info("Excel guardado: %s", EXCEL.name)

    problemas = verificar(respaldo, EXCEL, x14, columnas_editables)
    if problemas:
        log.error("La verificación encontró cambios fuera de los resultados:\n  %s", "\n  ".join(problemas))
    else:
        log.info("Verificación OK: fórmulas, otras pestañas y validaciones sin cambios")


# ---------------------------------------------------------------- reporte

def imprimir_tabla(filas, omitidos, modelos):
    print("\nResumen por modelo")
    print(f"{'Modelo':<17} {'Casos':>5} {'JSON válido':>11} {'Puntaje / n':>12} {'%':>6} {'Latencia media':>15} {'Errores':>7}")
    for m in modelos:
        if m is None:
            continue
        if m in omitidos:
            print(f"{m:<17} {omitidos[m]}")
            continue
        fs = [f for f in filas if f["modelo"] == m]
        if not fs:
            print(f"{m:<17} (sin resultados en esta corrida)")
            continue
        pct_json = 100 * sum(f["json_valido"] for f in fs) / len(fs)
        puntos, n = sum(f["puntaje"] for f in fs), sum(f["n"] for f in fs)
        lats = [f["latencia"] for f in fs if f["latencia"] is not None]
        lat = f"{sum(lats) / len(lats):.1f} s" if lats else "—"
        print(f"{m:<17} {len(fs):>5} {pct_json:>10.0f}% {f'{puntos:g} / {n}':>12} {100 * puntos / n:>5.0f}% "
              f"{lat:>15} {sum(bool(f['error']) for f in fs):>7}")

    ceros = defaultdict(list)
    for f in filas:
        if f["puntaje"] == 0:
            ceros[f["caso"]].append(f["modelo"] + (" (error)" if f["error"] else ""))
    if ceros:
        print("\nCasos con puntaje 0 (más fallados primero)")
        for caso, ms in sorted(ceros.items(), key=lambda x: (-len(x[1]), x[0])):
            print(f"  {caso}: {len(ms)} modelo(s) — {', '.join(ms)}")


# ---------------------------------------------------------------- main

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--casos", nargs="+", help="IDs de casos a correr (por defecto, los 25)")
    ap.add_argument("--modelos", nargs="+", help="Modelos a correr (por defecto, Evaluación!A5:A11)")
    ap.add_argument("--sin-excel", action="store_true", help="No escribir el Excel")
    ap.add_argument("--desde", type=Path, help="Llenar el Excel con una corrida guardada, sin llamar al backend")
    args = ap.parse_args()

    RESULTADOS.mkdir(exist_ok=True)
    stamp = datetime.now().strftime("%Y-%m-%d_%H%M%S")
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)-7s %(message)s", datefmt="%H:%M:%S",
                        handlers=[logging.StreamHandler(sys.stdout),
                                  logging.FileHandler(RESULTADOS / f"{stamp}.log", encoding="utf-8")])

    casos = cargar_casos()
    wb, _ = abrir_excel(EXCEL)
    modelos_excel = leer_modelos(wb)
    if args.modelos:
        desconocidos = set(args.modelos) - set(modelos_excel)
        if desconocidos:
            sys.exit(f"Estos modelos no están en Evaluación!A5:A11: {', '.join(sorted(desconocidos))}")

    if args.desde:
        filas, omitidos = cargar_corrida(args.desde, casos)
    else:
        if args.casos:
            casos = [c for c in casos if c["id"] in set(args.casos)]
        modelos = [m for m in modelos_excel if m and (not args.modelos or m in args.modelos)]
        prompts = cargar_prompts(casos)
        salida = RESULTADOS / f"{stamp}.jsonl"
        log.info("Corrida: %d modelo(s) × %d caso(s) → %s", len(modelos), len(casos), salida.relative_to(PROYECTO))
        try:
            filas, omitidos = correr(casos, prompts, modelos, salida)
        except SinEsquema as e:
            log.error("El backend no encuentra eval/esquemas.json: %s", e)
            sys.exit(f"\nReinicia el backend con RUTA_ESQUEMAS={EVAL / 'esquemas.json'}")

    if not args.sin_excel:
        escribir_excel(filas, omitidos, modelos_excel, stamp)
    imprimir_tabla(filas, omitidos, modelos_excel)


if __name__ == "__main__":
    main()
